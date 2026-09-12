from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "documents"
TMP_DIR = ROOT / "tmp" / "visual_playbook" / "pages"
OUTPUT = OUT_DIR / "MEC_Finance_Playbook_Visuel_2026_FR.docx"
LOGO = ROOT / "public" / "images" / "brand" / "logo-main.png"
HERO = ROOT / "assets" / "marketing_visuals" / "mec-advisory-meeting.png"
REEL = ROOT / "assets" / "marketing_visuals" / "mec-reel-expert.png"

W, H = 1600, 1212
NAVY = (0, 31, 63)
NAVY_2 = (8, 48, 85)
GOLD = (212, 175, 55)
GOLD_LIGHT = (239, 219, 151)
WHITE = (255, 255, 255)
INK = (20, 30, 42)
GRAY = (92, 103, 116)
LIGHT = (246, 248, 251)
PALE = (233, 240, 247)
BORDER = (214, 222, 230)
GREEN = (33, 132, 102)
RED = (178, 61, 61)

FONT_REG = Path("C:/Windows/Fonts/arial.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/arialbd.ttf")
FONT_ITALIC = Path("C:/Windows/Fonts/ariali.ttf")


def fnt(size: int, bold: bool = False, italic: bool = False):
    path = FONT_BOLD if bold else FONT_ITALIC if italic else FONT_REG
    return ImageFont.truetype(str(path), size)


def fit_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, start: int, minimum: int = 16, bold: bool = False):
    for size in range(start, minimum - 1, -1):
        font = fnt(size, bold=bold)
        if draw.textbbox((0, 0), text, font=font)[2] <= max_width:
            return font
    return fnt(minimum, bold=bold)


def wrap(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        line = words[0]
        for word in words[1:]:
            probe = f"{line} {word}"
            if draw.textbbox((0, 0), probe, font=font)[2] <= max_width:
                line = probe
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


def text_block(draw, xy, text, font, fill, max_width, spacing=8, anchor="la"):
    x, y = xy
    lines = wrap(draw, text, font, max_width)
    ascent, descent = font.getmetrics()
    line_h = ascent + descent + spacing
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill, anchor=anchor)
        y += line_h
    return y


def rounded(draw, box, radius=24, fill=WHITE, outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def crop_image(path: Path, size: tuple[int, int], focus=(0.5, 0.5)):
    image = Image.open(path).convert("RGB")
    return ImageOps.fit(image, size, method=Image.Resampling.LANCZOS, centering=focus)


def paste_rounded(canvas, image, box, radius=28):
    x1, y1, x2, y2 = box
    fitted = ImageOps.fit(image.convert("RGB"), (x2 - x1, y2 - y1), method=Image.Resampling.LANCZOS)
    mask = Image.new("L", fitted.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, fitted.width, fitted.height), radius=radius, fill=255)
    canvas.paste(fitted, (x1, y1), mask)


def base(title: str, kicker: str, subtitle: str = "", page: int = 1):
    img = Image.new("RGB", (W, H), LIGHT)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, W, 16), fill=GOLD)
    draw.text((70, 48), kicker.upper(), font=fnt(18, bold=True), fill=GOLD)
    draw.text((70, 84), title, font=fit_text(draw, title, 1320, 43, 28, bold=True), fill=INK)
    if subtitle:
        text_block(draw, (70, 145), subtitle, fnt(20), GRAY, 1260, 5)
    draw.text((1465, 55), f"{page:02d}", font=fnt(18, bold=True), fill=NAVY)
    draw.text((70, 1172), "MEC FINANCE   PLAYBOOK VISUEL 2026", font=fnt(14, bold=True), fill=(120, 130, 142))
    draw.rectangle((1270, 1179, 1530, 1183), fill=GOLD)
    return img, draw


def bullet(draw, x, y, text, max_width, color=INK, size=23, accent=GOLD):
    draw.ellipse((x, y + 9, x + 12, y + 21), fill=accent)
    return text_block(draw, (x + 28, y), text, fnt(size), color, max_width - 28, 7)


def arrow(draw, start, end, color=GOLD, width=7):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    if x2 >= x1:
        draw.polygon([(x2, y2), (x2 - 18, y2 - 11), (x2 - 18, y2 + 11)], fill=color)
    else:
        draw.polygon([(x2, y2), (x2 + 18, y2 - 11), (x2 + 18, y2 + 11)], fill=color)


def add_logo(canvas, box, white_card=False):
    logo = Image.open(LOGO).convert("RGBA")
    rgb = logo.convert("RGB")
    diff = ImageChops.difference(rgb, Image.new("RGB", rgb.size, WHITE)).convert("L")
    diff = diff.point(lambda value: 255 if value > 18 else 0)
    content_box = diff.getbbox()
    if content_box:
        logo = logo.crop(content_box)
    x1, y1, x2, y2 = box
    if white_card:
        overlay = ImageDraw.Draw(canvas)
        rounded(overlay, box, radius=22, fill=WHITE)
        pad = 18
        x1 += pad
        y1 += pad
        x2 -= pad
        y2 -= pad
    fitted = ImageOps.contain(logo, (x2 - x1, y2 - y1), method=Image.Resampling.LANCZOS)
    px = x1 + (x2 - x1 - fitted.width) // 2
    py = y1 + (y2 - y1 - fitted.height) // 2
    canvas.paste(fitted, (px, py), fitted)


def page_cover():
    hero = crop_image(HERO, (W, H), focus=(0.5, 0.48))
    hero = hero.filter(ImageFilter.GaussianBlur(0.2))
    img = hero.copy()
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    od.rectangle((0, 0, 790, H), fill=(*NAVY, 242))
    od.rectangle((790, 0, 1120, H), fill=(*NAVY, 105))
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    draw = ImageDraw.Draw(img)
    add_logo(img, (74, 70, 445, 218), white_card=True)
    draw.text((74, 350), "PLAYBOOK", font=fnt(23, bold=True), fill=GOLD)
    y = text_block(draw, (74, 400), "Stratégie marketing\net templates visuels", fnt(58, bold=True), WHITE, 650, 12)
    y += 24
    text_block(draw, (74, y), "Une version courte, présentable et directement exploitable pour LinkedIn, Instagram et Facebook.", fnt(24), (224, 232, 240), 585, 8)
    draw.text((74, 1080), "MEC FINANCE   MAROC  EUROPE   2026", font=fnt(18, bold=True), fill=WHITE)
    draw.rectangle((74, 1120, 450, 1128), fill=GOLD)
    return img


def page_decision():
    img, draw = base("La décision stratégique", "01  Positionnement", "La marque doit occuper une idée simple avant de multiplier les formats.", 2)
    rounded(draw, (70, 245, 1530, 515), radius=34, fill=NAVY)
    draw.text((115, 285), "CATÉGORIE À OCCUPER", font=fnt(19, bold=True), fill=GOLD)
    text_block(draw, (115, 340), "Le partenaire financier et opérationnel des entreprises qui travaillent entre le Maroc et l'Europe", fnt(42, bold=True), WHITE, 1330, 10)
    moves = [
        ("1", "INCARNER", "Un expert visible\nchaque semaine"),
        ("2", "EXPLIQUER", "Des réponses courtes\nà des problèmes réels"),
        ("3", "CONVERTIR", "Une ressource ou\nun diagnostic par sujet"),
    ]
    for i, (num, head, copy) in enumerate(moves):
        x = 70 + i * 500
        rounded(draw, (x, 575, x + 460, 1030), radius=28, fill=WHITE, outline=BORDER, width=2)
        draw.ellipse((x + 35, 615, x + 115, 695), fill=GOLD)
        draw.text((x + 75, 655), num, font=fnt(36, bold=True), fill=NAVY, anchor="mm")
        draw.text((x + 35, 745), head, font=fnt(25, bold=True), fill=NAVY)
        text_block(draw, (x + 35, 800), copy, fnt(29, bold=True), INK, 385, 12)
    return img


def page_competitors():
    img, draw = base("Ce que le marché nous apprend", "02  Benchmark", "Copier les mécanismes utiles, pas l'apparence des grandes marques.", 3)
    headers = [("UPSILON", "Régularité + ressources"), ("FAIR UP", "Positionnement Maroc France"), ("BDO", "Séries + vidéo"), ("PWC", "Experts + preuves terrain"), ("MCCG", "Voix du fondateur"), ("CABINETS LOCAUX", "Espace libre pour MEC")]
    for i, (name, lesson) in enumerate(headers):
        col = i % 3
        row = i // 3
        x = 70 + col * 500
        y = 235 + row * 290
        rounded(draw, (x, y, x + 460, y + 245), radius=26, fill=WHITE, outline=BORDER)
        draw.rectangle((x, y, x + 12, y + 245), fill=GOLD if i < 5 else NAVY)
        draw.text((x + 35, y + 34), name, font=fnt(21, bold=True), fill=NAVY)
        text_block(draw, (x + 35, y + 84), lesson, fnt(27, bold=True), INK, 385, 9)
    rounded(draw, (70, 855, 1530, 1055), radius=24, fill=PALE)
    draw.text((105, 890), "OPPORTUNITÉ MEC", font=fnt(20, bold=True), fill=GOLD)
    text_block(draw, (105, 938), "Plus de contenu pratique, plus de visages, plus de preuves et une signature éditoriale Maroc Europe cohérente.", fnt(30, bold=True), NAVY, 1340, 8)
    return img


def page_audiences():
    img, draw = base("Quatre publics à servir", "03  Audiences", "Chaque public reçoit un problème, un format et une prochaine étape spécifiques.", 4)
    cards = [
        ("01", "ENTREPRISE EUROPÉENNE", "S'implanter au Maroc", "Guide implantation", "Diagnostic projet"),
        ("02", "ENTREPRISE MAROCAINE", "Travailler avec l'Europe", "Checklist flux", "Revue fiscale"),
        ("03", "PME DE TANGER", "Piloter la performance", "Tableau de bord", "Diagnostic finance"),
        ("04", "INVESTISSEUR", "Opérer en zone franche", "Parcours réel", "Étude de faisabilité"),
    ]
    for i, (num, who, need, asset, cta) in enumerate(cards):
        col, row = i % 2, i // 2
        x, y = 70 + col * 750, 235 + row * 410
        rounded(draw, (x, y, x + 705, y + 365), radius=28, fill=WHITE, outline=BORDER)
        draw.ellipse((x + 35, y + 35, x + 105, y + 105), fill=NAVY)
        draw.text((x + 70, y + 70), num, font=fnt(24, bold=True), fill=WHITE, anchor="mm")
        draw.text((x + 135, y + 42), who, font=fnt(20, bold=True), fill=GOLD)
        draw.text((x + 35, y + 135), need, font=fnt(31, bold=True), fill=INK)
        draw.text((x + 35, y + 222), "CONTENU", font=fnt(16, bold=True), fill=GRAY)
        draw.text((x + 150, y + 218), asset, font=fnt(23), fill=NAVY)
        draw.text((x + 35, y + 285), "ACTION", font=fnt(16, bold=True), fill=GRAY)
        draw.text((x + 150, y + 281), cta, font=fnt(23, bold=True), fill=NAVY)
    return img


def page_channels():
    img, draw = base("Le système des canaux", "04  Distribution", "Un canal central, plusieurs usages et un parcours de conversion lisible.", 5)
    nodes = {
        "linkedin": (170, 455, 480, 655),
        "site": (650, 285, 950, 485),
        "whatsapp": (1120, 455, 1430, 655),
        "social": (650, 690, 950, 890),
    }
    arrow(draw, (480, 555), (650, 385), GOLD)
    arrow(draw, (950, 385), (1120, 555), GOLD)
    arrow(draw, (480, 590), (650, 790), NAVY)
    arrow(draw, (950, 790), (1120, 610), NAVY)
    content = [
        (nodes["linkedin"], NAVY, WHITE, "LINKEDIN", "Autorité\nExperts\nDocuments"),
        (nodes["site"], GOLD, NAVY, "SITE", "Articles\nGuides\nLanding pages"),
        (nodes["whatsapp"], GREEN, WHITE, "WHATSAPP", "Qualification\nRendez-vous\nSuivi"),
        (nodes["social"], WHITE, NAVY, "IG + FACEBOOK", "Reels\nProximité\nReciblage"),
    ]
    for box, fill, color, head, body in content:
        rounded(draw, box, radius=35, fill=fill, outline=BORDER if fill == WHITE else None)
        x1, y1, x2, y2 = box
        draw.text(((x1 + x2) // 2, y1 + 48), head, font=fnt(23, bold=True), fill=color, anchor="mm")
        text_block(draw, ((x1 + x2) // 2, y1 + 88), body, fnt(23, bold=True), color, x2 - x1 - 50, 6, anchor="ma")
    draw.text((800, 965), "YOUTUBE = ARCHIVE    •    TIKTOK = TEST AU MOIS 3", font=fnt(22, bold=True), fill=GRAY, anchor="mm")
    return img


def page_content_system():
    img, draw = base("Le moteur éditorial", "05  Contenu", "Quatre publications par semaine et six piliers clairement pondérés.", 6)
    # Donut chart
    cx, cy = 335, 600
    bbox = (115, 380, 555, 820)
    parts = [(25, GOLD), (20, NAVY), (20, (64, 103, 145)), (15, (196, 155, 52)), (10, (130, 151, 175)), (10, (185, 196, 207))]
    start = -90
    for value, color in parts:
        end = start + value * 3.6
        draw.pieslice(bbox, start=start, end=end, fill=color)
        start = end
    draw.ellipse((220, 485, 450, 715), fill=LIGHT)
    draw.text((335, 565), "4", font=fnt(66, bold=True), fill=NAVY, anchor="mm")
    draw.text((335, 635), "POSTS / SEMAINE", font=fnt(18, bold=True), fill=GRAY, anchor="mm")
    labels = [
        ("25%", "Maroc Europe", GOLD), ("20%", "Fiscalité", NAVY), ("20%", "Pilotage", (64, 103, 145)),
        ("15%", "Implantation", (196, 155, 52)), ("10%", "Preuves", (130, 151, 175)), ("10%", "Équipe", (185, 196, 207)),
    ]
    for i, (pct, lab, color) in enumerate(labels):
        x = 650 + (i % 2) * 430
        y = 275 + (i // 2) * 125
        draw.rounded_rectangle((x, y, x + 390, y + 95), radius=20, fill=WHITE, outline=BORDER)
        draw.rectangle((x, y, x + 14, y + 95), fill=color)
        draw.text((x + 35, y + 26), pct, font=fnt(27, bold=True), fill=NAVY)
        draw.text((x + 125, y + 29), lab, font=fnt(22, bold=True), fill=INK)
    rhythm = [("LUN", "CARROUSEL"), ("MAR", "REEL"), ("JEU", "CAS"), ("VEN", "CONFIANCE")]
    for i, (day, form) in enumerate(rhythm):
        x = 650 + i * 215
        y = 765
        rounded(draw, (x, y, x + 190, y + 205), radius=24, fill=NAVY if i == 1 else WHITE, outline=BORDER)
        draw.text((x + 95, y + 48), day, font=fnt(20, bold=True), fill=GOLD, anchor="mm")
        text_block(draw, (x + 95, y + 95), form, fnt(20, bold=True), WHITE if i == 1 else NAVY, 155, 5, anchor="ma")
    return img


def square_slide(size=430, bg=WHITE):
    return Image.new("RGB", (size, size), bg)


def slide_brand(slide, number=None, dark=False):
    d = ImageDraw.Draw(slide)
    color = WHITE if dark else NAVY
    d.rectangle((0, 0, slide.width, 12), fill=GOLD)
    d.text((24, 24), "MEC FINANCE", font=fnt(17, bold=True), fill=color)
    if number is not None:
        d.text((slide.width - 28, slide.height - 28), f"{number}/7", font=fnt(14, bold=True), fill=color, anchor="rs")


def page_carousel_rules():
    img, draw = base("Template Carousel", "06  Design", "Un système carré, lisible sur mobile et reconnaissable sans surcharge.", 7)
    # Cover example
    sample = square_slide(650, NAVY)
    sd = ImageDraw.Draw(sample)
    slide_brand(sample, 1, dark=True)
    sd.text((45, 115), "IMPLANTATION", font=fnt(22, bold=True), fill=GOLD)
    text_block(sd, (45, 165), "Créer une société\nau Maroc depuis\nl'Europe", fnt(47, bold=True), WHITE, 555, 10)
    sd.text((45, 475), "7 vérifications avant de signer", font=fnt(24, bold=True), fill=GOLD_LIGHT)
    sd.rounded_rectangle((45, 545, 300, 600), radius=18, fill=GOLD)
    sd.text((172, 573), "ENREGISTRER", font=fnt(18, bold=True), fill=NAVY, anchor="mm")
    img.paste(sample, (70, 245))
    rules = [
        ("1080 × 1080", "Format carré"),
        ("7 pages", "Une idée par page"),
        ("40 à 65 pt", "Titres courts"),
        ("1 CTA", "Dernière page"),
        ("Logo discret", "Même position"),
        ("80 px", "Zone de sécurité"),
    ]
    for i, (value, label) in enumerate(rules):
        x = 805 + (i % 2) * 360
        y = 260 + (i // 2) * 210
        rounded(draw, (x, y, x + 330, y + 170), radius=24, fill=WHITE, outline=BORDER)
        draw.text((x + 28, y + 33), value, font=fnt(28, bold=True), fill=NAVY)
        draw.text((x + 28, y + 93), label, font=fnt(20), fill=GRAY)
    rounded(draw, (805, 905, 1530, 1035), radius=22, fill=PALE)
    text_block(draw, (835, 930), "Bleu pour l'autorité. Doré pour guider l'œil. Blanc pour laisser respirer l'information.", fnt(23, bold=True), NAVY, 655, 8)
    return img


def make_carousel_slides():
    slides = []
    # 1 cover
    s = square_slide(420, NAVY); d = ImageDraw.Draw(s); slide_brand(s, 1, True)
    d.text((28, 92), "IMPLANTATION", font=fnt(17, bold=True), fill=GOLD)
    text_block(d, (28, 130), "Créer une société\nau Maroc depuis\nl'Europe", fnt(34, bold=True), WHITE, 360, 5)
    d.text((28, 327), "7 vérifications avant de signer", font=fnt(17, bold=True), fill=GOLD_LIGHT)
    slides.append(s)
    # 2
    s = square_slide(420, WHITE); d = ImageDraw.Draw(s); slide_brand(s, 2)
    d.text((28, 95), "01", font=fnt(65, bold=True), fill=GOLD)
    d.text((28, 175), "FORME JURIDIQUE", font=fnt(21, bold=True), fill=NAVY)
    text_block(d, (28, 225), "La gouvernance doit guider le choix.", fnt(27, bold=True), INK, 350, 6)
    slides.append(s)
    # 3 photo
    photo = crop_image(HERO, (420, 420), focus=(0.60, 0.55)); s = photo.copy(); d = ImageDraw.Draw(s)
    overlay = Image.new("RGBA", (420, 420), (*NAVY, 185)); s = Image.alpha_composite(s.convert("RGBA"), overlay).convert("RGB"); d = ImageDraw.Draw(s); slide_brand(s, 3, True)
    d.text((28, 100), "02  BANQUE", font=fnt(22, bold=True), fill=GOLD)
    text_block(d, (28, 155), "Préparez activité, bénéficiaires et flux.", fnt(31, bold=True), WHITE, 345, 8)
    slides.append(s)
    # 4
    s = square_slide(420, PALE); d = ImageDraw.Draw(s); slide_brand(s, 4)
    d.text((28, 95), "03", font=fnt(65, bold=True), fill=GOLD)
    d.text((28, 175), "FISCALITÉ", font=fnt(21, bold=True), fill=NAVY)
    text_block(d, (28, 225), "TVA, retenues et conventions avant la première facture.", fnt(25, bold=True), INK, 350, 6)
    slides.append(s)
    # 5
    s = square_slide(420, WHITE); d = ImageDraw.Draw(s); slide_brand(s, 5)
    d.text((28, 95), "04", font=fnt(65, bold=True), fill=GOLD)
    d.text((28, 175), "PREMIER SALARIÉ", font=fnt(21, bold=True), fill=NAVY)
    text_block(d, (28, 225), "Anticipez paie et CNSS dès le départ.", fnt(27, bold=True), INK, 350, 6)
    slides.append(s)
    # 6 checklist
    s = square_slide(420, NAVY); d = ImageDraw.Draw(s); slide_brand(s, 6, True)
    d.text((28, 92), "CHECKLIST", font=fnt(23, bold=True), fill=GOLD)
    items = ["Structure", "Banque", "Fiscalité", "RH", "Reporting"]
    for i, item in enumerate(items):
        y = 145 + i * 49
        d.rounded_rectangle((30, y, 56, y + 26), radius=6, outline=GOLD, width=3)
        d.text((75, y - 1), item, font=fnt(22, bold=True), fill=WHITE)
    slides.append(s)
    # 7 CTA
    s = square_slide(420, GOLD); d = ImageDraw.Draw(s); slide_brand(s, 7)
    d.text((28, 103), "PRÊT À AVANCER", font=fnt(20, bold=True), fill=NAVY)
    text_block(d, (28, 155), "Téléchargez la checklist", fnt(37, bold=True), NAVY, 350, 8)
    d.rounded_rectangle((28, 292, 365, 352), radius=18, fill=NAVY)
    d.text((196, 322), "mecfinance.ma", font=fnt(23, bold=True), fill=WHITE, anchor="mm")
    slides.append(s)
    return slides


def page_carousel_storyboard():
    img, draw = base("Exemple Carousel en 7 pages", "07  Template prêt à adapter", "Sujet exemple  Créer une société au Maroc depuis l'Europe", 8)
    slides = make_carousel_slides()
    thumb = 320
    top_y = 260
    positions = []
    for i in range(4):
        positions.append((70 + i * 380, top_y))
    for i in range(3):
        positions.append((260 + i * 380, top_y + 405))
    for slide, (x, y) in zip(slides, positions):
        slide = slide.resize((thumb, thumb), Image.Resampling.LANCZOS)
        shadow = Image.new("RGBA", (thumb + 20, thumb + 20), (0, 0, 0, 0))
        ImageDraw.Draw(shadow).rounded_rectangle((10, 10, thumb + 10, thumb + 10), radius=18, fill=(0, 0, 0, 35))
        img.paste(shadow, (x - 10, y - 4), shadow)
        img.paste(slide, (x, y))
    return img


def reel_frame(photo: Image.Image, variant: int):
    frame = ImageOps.fit(photo.convert("RGB"), (300, 533), method=Image.Resampling.LANCZOS, centering=(0.50, 0.50))
    if variant == 1:
        overlay = Image.new("RGBA", frame.size, (0, 0, 0, 0)); od = ImageDraw.Draw(overlay); od.rectangle((0, 0, 300, 230), fill=(*NAVY, 205)); frame = Image.alpha_composite(frame.convert("RGBA"), overlay).convert("RGB")
        d = ImageDraw.Draw(frame); d.text((22, 28), "LA MINUTE MEC", font=fnt(14, bold=True), fill=GOLD)
        text_block(d, (22, 70), "Rentable…\nmais sans\ntrésorerie ?", fnt(30, bold=True), WHITE, 250, 4)
    elif variant == 2:
        overlay = Image.new("RGBA", frame.size, (0, 0, 0, 0)); od = ImageDraw.Draw(overlay); od.rectangle((0, 330, 300, 533), fill=(*NAVY, 215)); frame = Image.alpha_composite(frame.convert("RGBA"), overlay).convert("RGB")
        d = ImageDraw.Draw(frame); d.text((22, 352), "3 CHIFFRES", font=fnt(18, bold=True), fill=GOLD)
        text_block(d, (22, 392), "Délai client\nDélai fournisseur\nStock moyen", fnt(20, bold=True), WHITE, 250, 3)
    else:
        overlay = Image.new("RGBA", frame.size, (*NAVY, 230)); frame = Image.alpha_composite(frame.convert("RGBA"), overlay).convert("RGB")
        d = ImageDraw.Draw(frame); d.text((150, 120), "SUIVEZ LE BFR", font=fnt(22, bold=True), fill=GOLD, anchor="mm")
        text_block(d, (150, 180), "Chaque mois\npas seulement\nà la clôture", fnt(30, bold=True), WHITE, 250, 6, anchor="ma")
        d.rounded_rectangle((38, 405, 262, 470), radius=18, fill=GOLD)
        d.text((150, 438), "VOIR LE GUIDE", font=fnt(17, bold=True), fill=NAVY, anchor="mm")
    d = ImageDraw.Draw(frame)
    d.rectangle((0, 0, 300, 8), fill=GOLD)
    d.text((20, 500), "MEC FINANCE", font=fnt(13, bold=True), fill=WHITE)
    return frame


def page_reel_design():
    img, draw = base("Template Reel", "08  Design", "Le visage crée la confiance. Le texte guide la lecture sans son.", 9)
    photo = Image.open(REEL).convert("RGB")
    labels = [("0–3 s", "HOOK"), ("3–30 s", "VALEUR"), ("30–45 s", "CTA")]
    for i, (timing, label) in enumerate(labels):
        x = 70 + i * 390
        y = 250
        frame = reel_frame(photo, i + 1)
        img.paste(frame, (x, y))
        draw.text((x + 150, y + 575), timing, font=fnt(24, bold=True), fill=NAVY, anchor="mm")
        draw.text((x + 150, y + 620), label, font=fnt(17, bold=True), fill=GOLD, anchor="mm")
    rounded(draw, (1260, 250, 1530, 930), radius=28, fill=NAVY)
    draw.text((1395, 300), "RÈGLES", font=fnt(22, bold=True), fill=GOLD, anchor="mm")
    rules = ["9:16", "30–60 s", "Sous-titres", "1 question", "1 CTA", "Logo à la fin"]
    for i, rule in enumerate(rules):
        y = 370 + i * 82
        draw.ellipse((1300, y + 4, 1316, y + 20), fill=GOLD)
        draw.text((1335, y), rule, font=fnt(21, bold=True), fill=WHITE)
    draw.text((800, 1055), "Même vidéo  →  LinkedIn  •  Instagram  •  Facebook  •  YouTube Shorts", font=fnt(24, bold=True), fill=GRAY, anchor="mm")
    return img


def page_reel_script():
    img, draw = base("Reel prêt à tourner", "09  Storyboard", "Sujet exemple  Pourquoi le bénéfice ne garantit pas la trésorerie", 10)
    segments = [
        ("0–3 s", "QUESTION", "Votre société est rentable. Pourquoi manque-t-elle toujours de trésorerie ?", GOLD),
        ("3–12 s", "CONSTAT", "Le bénéfice ne dit pas quand les clients paient ni quand les charges sortent.", NAVY),
        ("12–30 s", "EXPLICATION", "Regardez trois chiffres : délai client, délai fournisseur et stock moyen.", (64, 103, 145)),
        ("30–40 s", "DÉCISION", "Si le BFR augmente plus vite que le chiffre d'affaires, la croissance consomme du cash.", GREEN),
        ("40–45 s", "ACTION", "Enregistrez la vidéo ou ouvrez le guide BFR.", GOLD),
    ]
    y = 245
    for timing, phase, copy, color in segments:
        rounded(draw, (70, y, 1530, y + 145), radius=24, fill=WHITE, outline=BORDER)
        draw.rounded_rectangle((90, y + 30, 245, y + 112), radius=20, fill=color)
        draw.text((167, y + 71), timing, font=fnt(22, bold=True), fill=WHITE if color != GOLD else NAVY, anchor="mm")
        draw.text((285, y + 28), phase, font=fnt(18, bold=True), fill=GOLD)
        text_block(draw, (285, y + 66), copy, fnt(24, bold=True), INK, 1170, 5)
        y += 165
    draw.text((70, 1093), "TOURNAGE  plan poitrine  •  regard caméra  •  lumière naturelle  •  sous-titres mot à mot", font=fnt(20, bold=True), fill=NAVY)
    return img


def page_calendar():
    img, draw = base("Plan de contenu des 30 premiers jours", "10  Calendrier", "Un mois pour installer la signature visuelle et mesurer les premiers signaux.", 11)
    weeks = [
        ("S1", ["Implantation\nCarrousel", "Qui est MEC\nReel", "Équipe\nPhoto", "3 chiffres\nPost AR"]),
        ("S2", ["SARL ou SAS\nCarrousel", "TVA ≠ marge\nReel", "Échéance\nAlerte", "Quel chiffre ?\nQuestion"]),
        ("S3", ["Checklist\nDocument", "Compte bancaire\nReel", "Préparation\nPhoto", "Premier salarié\nPost"]),
        ("S4", ["Trésorerie\nCas", "Dashboard\nReel", "Échéance\nPost AR", "Méthode MEC\nExpert"]),
    ]
    days = ["LUNDI", "MARDI", "JEUDI", "VENDREDI"]
    x0, y0 = 180, 300
    cell_w, cell_h = 320, 170
    for c, day in enumerate(days):
        x = x0 + c * cell_w
        draw.rounded_rectangle((x + 5, 235, x + cell_w - 12, 285), radius=14, fill=NAVY)
        draw.text((x + cell_w / 2 - 4, 260), day, font=fnt(18, bold=True), fill=WHITE, anchor="mm")
    for r, (week, cells) in enumerate(weeks):
        y = y0 + r * cell_h
        draw.text((118, y + 70), week, font=fnt(24, bold=True), fill=GOLD, anchor="mm")
        for c, copy in enumerate(cells):
            x = x0 + c * cell_w
            fill = PALE if (r + c) % 2 else WHITE
            rounded(draw, (x + 5, y, x + cell_w - 12, y + 145), radius=20, fill=fill, outline=BORDER)
            text_block(draw, (x + 25, y + 30), copy, fnt(22, bold=True), NAVY, cell_w - 60, 4)
    rounded(draw, (180, 1005, 1450, 1085), radius=20, fill=NAVY)
    draw.text((815, 1045), "OBJECTIF DU MOIS  16 publications  •  4 Reels  •  1 ressource  •  1 base de mesure", font=fnt(22, bold=True), fill=WHITE, anchor="mm")
    return img


def page_roadmap():
    img, draw = base("Feuille de route sur 90 jours", "11  Exécution", "Construire la confiance avant d'augmenter les dépenses média.", 12)
    phases = [
        ("0–30", "FONDATIONS", GOLD, ["Comptes + site alignés", "20 photos + 4 vidéos", "3 templates Carousel", "16 publications"]),
        ("31–60", "PREUVES", NAVY, ["2 articles profonds", "1 landing page", "1 newsletter", "Experts actifs"]),
        ("61–90", "CONVERSION", GREEN, ["Scorecard", "Premier cas publié", "Tests langues", "Budget sur gagnants"]),
    ]
    for i, (days, title, color, items) in enumerate(phases):
        x = 70 + i * 500
        rounded(draw, (x, 265, x + 460, 975), radius=30, fill=WHITE, outline=BORDER)
        draw.rounded_rectangle((x, 265, x + 460, 390), radius=30, fill=color)
        draw.rectangle((x, 350, x + 460, 390), fill=color)
        draw.text((x + 35, 300), days, font=fnt(41, bold=True), fill=WHITE if color != GOLD else NAVY)
        draw.text((x + 35, 355), title, font=fnt(20, bold=True), fill=WHITE if color != GOLD else NAVY)
        y = 455
        for item in items:
            y = bullet(draw, x + 35, y, item, 385, size=23, accent=color) + 22
    draw.text((800, 1050), "NE PAS LANCER DE CAMPAGNE LARGE AVANT 4 SEMAINES DE CONTENU ORGANIQUE", font=fnt(22, bold=True), fill=RED, anchor="mm")
    return img


def page_kpis():
    img, draw = base("Le tableau de bord utile", "12  Mesure", "La bonne question n'est pas combien de likes, mais quelles décisions et quelles opportunités.", 13)
    levels = [
        ("VISIBILITÉ", "Portée ciblée\nVues 50%\nEnregistrements", (64, 103, 145), 1220),
        ("CONFIANCE", "Commentaires dirigeants\nTéléchargements\nMessages qualifiés", NAVY, 940),
        ("CONVERSION", "Diagnostics\nPropositions\nContrats", GOLD, 660),
    ]
    y = 270
    for title, copy, color, width in levels:
        x = (W - width) // 2
        draw.polygon([(x, y), (x + width, y), (x + width - 70, y + 190), (x + 70, y + 190)], fill=color)
        draw.text((x + width / 2, y + 48), title, font=fnt(24, bold=True), fill=WHITE if color != GOLD else NAVY, anchor="mm")
        text_block(draw, (x + width / 2, y + 88), copy, fnt(22, bold=True), WHITE if color != GOLD else NAVY, width - 180, 4, anchor="ma")
        y += 210
    rounded(draw, (180, 940, 1420, 1070), radius=24, fill=WHITE, outline=BORDER)
    draw.text((225, 972), "DÉCISION HEBDOMADAIRE", font=fnt(19, bold=True), fill=GOLD)
    draw.text((225, 1017), "Renforcer le sujet qui produit des conversations qualifiées, pas celui qui attire le plus de réactions.", font=fnt(25, bold=True), fill=NAVY)
    return img


def page_guardrails():
    img, draw = base("Règles de marque et de conformité", "13  Garde-fous", "Une communication premium reste précise, sobre, datée et juridiquement vérifiée.", 14)
    rounded(draw, (70, 245, 755, 940), radius=28, fill=WHITE, outline=BORDER)
    rounded(draw, (845, 245, 1530, 940), radius=28, fill=WHITE, outline=BORDER)
    draw.text((110, 295), "À FAIRE", font=fnt(27, bold=True), fill=GREEN)
    draw.text((885, 295), "À ÉVITER", font=fnt(27, bold=True), fill=RED)
    dos = ["Une question par contenu", "Sources et dates visibles", "Experts et lieux réels", "Cas anonymisés", "Un seul appel à l'action"]
    donts = ["Promesses garanties", "Images génériques répétées", "Cinq sujets par visuel", "Clients cités sans accord", "Publicité avant validation"]
    y = 370
    for item in dos:
        draw.ellipse((112, y + 5, 136, y + 29), fill=GREEN)
        draw.line((118, y + 17, 122, y + 22), fill=WHITE, width=3)
        draw.line((122, y + 22, 131, y + 12), fill=WHITE, width=3)
        draw.text((160, y), item, font=fnt(25, bold=True), fill=INK)
        y += 100
    y = 370
    for item in donts:
        draw.ellipse((887, y + 5, 911, y + 29), fill=RED)
        draw.text((899, y + 17), "×", font=fnt(18, bold=True), fill=WHITE, anchor="mm")
        draw.text((935, y), item, font=fnt(25, bold=True), fill=INK)
        y += 100
    draw.rounded_rectangle((70, 985, 1530, 1080), radius=22, fill=NAVY)
    draw.text((800, 1018), "VALIDATION AVANT DIFFUSION", font=fnt(18, bold=True), fill=GOLD, anchor="mm")
    draw.text((800, 1054), "Expert métier  •  conformité  •  identité  •  lien de suivi", font=fnt(24, bold=True), fill=WHITE, anchor="mm")
    return img


def save_pages():
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    pages = [
        page_cover(), page_decision(), page_competitors(), page_audiences(), page_channels(),
        page_content_system(), page_carousel_rules(), page_carousel_storyboard(), page_reel_design(),
        page_reel_script(), page_calendar(), page_roadmap(), page_kpis(), page_guardrails(),
    ]
    paths = []
    for i, page in enumerate(pages, 1):
        path = TMP_DIR / f"page-{i:02d}.png"
        page.save(path, quality=95)
        paths.append(path)
    return paths


def set_alt(picture, description: str, title: str):
    picture._inline.docPr.set("descr", description)
    picture._inline.docPr.set("title", title)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("MEC Finance  |  Playbook visuel 2026  |  ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120, 130, 142)
    field = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    field._r.extend([begin, instr, separate, text, end])
    field.font.name = "Arial"; field.font.size = Pt(8); field.font.color.rgb = RGBColor(120, 130, 142)


def build_docx(page_paths: Iterable[Path]):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.left_margin = Inches(0.35)
    sec.right_margin = Inches(0.35)
    sec.top_margin = Inches(0.26)
    sec.bottom_margin = Inches(0.43)
    sec.header_distance = Inches(0.1)
    sec.footer_distance = Inches(0.15)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)

    footer = sec.footer
    footer.paragraphs[0].clear()
    add_page_field(footer.paragraphs[0])

    alt_titles = [
        "Couverture du playbook visuel", "Décision stratégique", "Benchmark concurrentiel",
        "Publics prioritaires", "Système des canaux", "Moteur éditorial", "Template Carousel",
        "Exemple Carousel en sept pages", "Template Reel", "Storyboard Reel", "Calendrier de 30 jours",
        "Feuille de route de 90 jours", "Tableau de bord", "Règles de marque et conformité",
    ]
    for i, (path, alt_title) in enumerate(zip(page_paths, alt_titles)):
        if i:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        picture = p.add_run().add_picture(str(path), width=Inches(10.30), height=Inches(7.80))
        set_alt(picture, f"Page visuelle MEC Finance consacrée à {alt_title.lower()}", alt_title)

    doc.core_properties.title = "Stratégie marketing et templates visuels de MEC Finance"
    doc.core_properties.subject = "Playbook visuel pour les Carousels, Reels et contenus sociaux"
    doc.core_properties.author = "MEC Finance"
    doc.core_properties.keywords = "MEC Finance, marketing, social media, carousel, reel, Maroc, Europe"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx(save_pages())
