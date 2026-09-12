from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
LANG = "fr" if "--lang" in sys.argv and sys.argv[sys.argv.index("--lang") + 1].lower() == "fr" else "ar"
SOURCE = ROOT / ("report-source-fr.md" if LANG == "fr" else "report-source.md")
OUTPUT_DIR = ROOT / "output" / "documents"
OUTPUT = OUTPUT_DIR / ("MEC_Finance_Strategie_Marketing_2026_FR.docx" if LANG == "fr" else "MEC_Finance_Marketing_Strategy_2026.docx")
LOGO = ROOT / "public" / "images" / "brand" / "logo-main.png"

NAVY = "001F3F"
GOLD = "D4AF37"
PALE_BLUE = "F2F6FA"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = "404040"
WHITE = "FFFFFF"
BLACK = "000000"
FONT = "Arial"
URL_RE = re.compile(r"https?://[^\s)]+")
AR_RE = re.compile(r"[\u0600-\u06FF]")
LINK_LABEL = "Ouvrir la source" if LANG == "fr" else "فتح المصدر"


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_direction(paragraph, rtl=True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if rtl:
        if bidi is None:
            bidi = OxmlElement("w:bidi")
            p_pr.append(bidi)
        bidi.set(qn("w:val"), "1")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        if bidi is not None:
            p_pr.remove(bidi)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_hyperlink(paragraph, text, url, color=NAVY):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:cs"), FONT)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_el, underline])
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, size=11, color=TEXT_GRAY):
    pos = 0
    bold = False
    tokens = re.split(r"(\*\*|https?://[^\s)]+)", text)
    for token in tokens:
        if not token:
            continue
        if token == "**":
            bold = not bold
            continue
        if token.startswith("http://") or token.startswith("https://"):
            add_hyperlink(paragraph, LINK_LABEL, token)
            continue
        run = paragraph.add_run(token)
        set_run_font(run, size=size, bold=bold, color=color)
        if AR_RE.search(token):
            run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))
        pos += len(token)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color="777777")


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title._element.rPr.rFonts.set(qn("w:cs"), FONT)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name, size, before, after in (
        ("Heading 1", 19, 18, 8),
        ("Heading 2", 15, 14, 6),
        ("Heading 3", 12.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for sec in doc.sections:
        footer = sec.footer
        footer.is_linked_to_previous = True
        p = footer.paragraphs[0]
        p.text = ""
        footer_text = "MEC Finance  |  Étude concurrentielle et stratégie marketing  |  " if LANG == "fr" else "MEC Finance  |  Research and Marketing Strategy  |  "
        run = p.add_run(footer_text)
        set_run_font(run, size=8.5, color="777777")
        add_page_number(p)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        picture = p.add_run().add_picture(str(LOGO), width=Inches(4.4))
        picture._inline.docPr.set("descr", "Logo Morocco Europe Consulting" if LANG == "fr" else "شعار Morocco Europe Consulting")
        picture._inline.docPr.set("title", "MEC Finance")
    p.paragraph_format.space_after = Pt(38)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Stratégie marketing et plan de contenu de MEC Finance" if LANG == "fr" else "خطة التسويق والمحتوى الرقمي لشركة MEC Finance")
    for run in title.runs:
        set_run_font(run, size=26, bold=True, color=BLACK)
        run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))
    set_paragraph_direction(title, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Étude concurrentielle et plan d'exécution sur 90 jours" if LANG == "fr" else "بحث تنافسي واستراتيجية تنفيذ لمدة 90 يوما")
    set_run_font(run, size=15, bold=True, color=NAVY)
    run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))
    set_paragraph_direction(subtitle, True)
    subtitle.paragraph_format.space_before = Pt(8)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(meta, "Tanger, Maroc et marchés européens  |  8 septembre 2026" if LANG == "fr" else "طنجة والمغرب والأسواق الأوروبية  |  8 سبتمبر 2026", size=10.5, color="666666")
    meta.paragraph_format.space_before = Pt(22)
    set_paragraph_direction(meta, True)

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(statement, "Document de travail pour la direction, le marketing et le développement commercial" if LANG == "fr" else "وثيقة عمل للإدارة والتسويق وتطوير الأعمال", size=10.5, color=GOLD)
    set_paragraph_direction(statement, True)
    statement.paragraph_format.space_before = Pt(90)
    doc.add_page_break()


def add_toc(doc):
    h = doc.add_paragraph("Sommaire" if LANG == "fr" else "محتويات التقرير", style="Heading 1")
    set_paragraph_direction(h, LANG != "fr")
    items = ([
        "Synthèse décisionnelle",
        "Offre réelle de MEC",
        "État de la présence numérique",
        "Potentiel numérique au Maroc",
        "Cartographie des concurrents",
        "Pratiques qui fonctionnent",
        "Positionnement et publics prioritaires",
        "Canaux et système de contenu",
        "Calendrier éditorial sur 12 semaines",
        "Modèles prêts à publier",
        "Conversion, publicité et mesure",
        "Feuille de route, risques et conformité",
        "Sources et limites de l'étude",
    ] if LANG == "fr" else [
        "الجواب التنفيذي",
        "ما الذي تبيعه MEC فعليا",
        "وضع الحضور الرقمي الحالي",
        "حجم الفرصة الرقمية في المغرب",
        "خريطة المنافسين",
        "ما الذي ينجح عند المنافسين",
        "التموضع والجماهير ذات الأولوية",
        "استراتيجية القنوات ونظام المحتوى",
        "تقويم النشر لمدة 12 أسبوعا",
        "نماذج جاهزة للنشر",
        "مسار التحويل والإعلان والقياس",
        "خارطة التنفيذ والمخاطر والامتثال",
        "المصادر وحدود البحث",
    ])
    for i, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        add_inline(p, f"{i:02d}  {item}", size=11.5, color=TEXT_GRAY)
        set_paragraph_direction(p, LANG != "fr")
    doc.add_page_break()


def clean_heading(text):
    return re.sub(r"[*_`]", "", text).strip()


def is_rtl(text):
    arabic = len(AR_RE.findall(text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return arabic >= max(1, latin // 3)


def add_table(doc, rows):
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])

    if len(headers) == 6:
        widths = [1.25, 0.55, 0.65, 0.75, 0.75, 2.35]
    elif len(headers) == 5:
        widths = [1.0, 1.35, 1.35, 1.1, 2.0]
    elif len(headers) == 4:
        widths = [1.25, 1.0, 2.0, 2.5]
    elif len(headers) == 3:
        widths = [0.95, 3.7, 1.95]
    else:
        widths = [6.7 / len(headers)] * len(headers)

    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.text = ""
        shade_cell(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell, top=95, start=95, bottom=95, end=95)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        add_inline(p, text.strip(), size=9.2, color=WHITE)
        for run in p.runs:
            run.bold = True
        set_paragraph_direction(p, is_rtl(text))

    for row_index, data in enumerate(rows[1:], start=1):
        row = table.add_row()
        for idx, text in enumerate(data):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            cell.text = ""
            if row_index % 2 == 0:
                shade_cell(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell, top=90, start=95, bottom=90, end=95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            add_inline(p, text.strip(), size=8.9, color=TEXT_GRAY)
            if idx > 0 and len(text.strip()) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                set_paragraph_direction(p, is_rtl(text))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_toc(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    table_rows = []
    skip_front_title = True
    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()

        if stripped == "## سجل الادعاءات والمصادر":
            break

        if stripped.startswith("# ") and skip_front_title:
            index += 1
            continue
        if stripped.startswith("## بحث تنافسي"):
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                index += 1
                continue
            table_rows.append(cells)
            next_line = lines[index + 1].strip() if index + 1 < len(lines) else ""
            if not (next_line.startswith("|") and next_line.endswith("|")):
                add_table(doc, table_rows)
                table_rows = []
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("### "):
            text = clean_heading(stripped[4:])
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, text, size=15, color=BLACK)
            for run in p.runs:
                run.bold = True
            set_paragraph_direction(p, is_rtl(text))
        elif stripped.startswith("## "):
            text = clean_heading(stripped[3:])
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, text, size=19, color=BLACK)
            for run in p.runs:
                run.bold = True
            set_paragraph_direction(p, is_rtl(text))
        elif stripped.startswith("# "):
            text = clean_heading(stripped[2:])
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, text, size=19, color=BLACK)
            for run in p.runs:
                run.bold = True
            set_paragraph_direction(p, is_rtl(text))
        elif re.match(r"^- ", stripped):
            text = stripped[2:]
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text, size=10.8)
            set_paragraph_direction(p, is_rtl(text))
        elif re.match(r"^\d+\. ", stripped):
            text = re.sub(r"^\d+\. ", "", stripped)
            p = doc.add_paragraph(style="List Number")
            add_inline(p, text, size=10.8)
            set_paragraph_direction(p, is_rtl(text))
        else:
            paragraph_text = stripped
            while index + 1 < len(lines):
                nxt = lines[index + 1].strip()
                if not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("- ") or re.match(r"^\d+\. ", nxt):
                    break
                paragraph_text += " " + nxt
                index += 1
            p = doc.add_paragraph()
            add_inline(p, paragraph_text, size=11)
            set_paragraph_direction(p, is_rtl(paragraph_text))
            if paragraph_text.startswith("**") and paragraph_text.endswith("**"):
                p.paragraph_format.keep_with_next = True

        index += 1

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    doc.core_properties.title = "Stratégie marketing et plan de contenu de MEC Finance" if LANG == "fr" else "خطة التسويق والمحتوى الرقمي لشركة MEC Finance"
    doc.core_properties.subject = "Étude concurrentielle et stratégie de contenu sur 90 jours" if LANG == "fr" else "بحث تنافسي واستراتيجية محتوى لمدة 90 يوما"
    doc.core_properties.author = "MEC Finance"
    doc.core_properties.keywords = "MEC Finance, marketing, social media, Morocco, Europe"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
